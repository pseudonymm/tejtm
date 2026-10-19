from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from tejtm.models import SourceColumn, SourceProfile, TargetField, TargetSchema

SCENARIOS_ROOT = Path(__file__).resolve().parent.parent / "fixtures" / "scenarios"


def list_scenario_dirs() -> list[Path]:
    return sorted(p for p in SCENARIOS_ROOT.iterdir() if p.is_dir() and (p / "scenario.json").exists())


def load_scenario(path: Path) -> dict[str, Any]:
    return json.loads((path / "scenario.json").read_text(encoding="utf-8"))


def _read_csv_columns(data_path: Path, sample_rows: int) -> list[SourceColumn]:
    with data_path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames or []
        buckets: dict[str, list[str]] = {name: [] for name in fieldnames}
        for i, row in enumerate(reader):
            if i >= sample_rows:
                break
            for name in fieldnames:
                val = row.get(name)
                if val is not None and str(val).strip():
                    buckets[name].append(str(val))
    return [SourceColumn(name=name, samples=tuple(buckets[name])) for name in fieldnames]


def _read_xlsx_columns(data_path: Path, sample_rows: int) -> list[SourceColumn]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("openpyxl required for xlsx scenario tests") from exc
    wb = load_workbook(data_path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    names = [str(h) if h is not None else "" for h in header]
    buckets: dict[str, list[str]] = {n: [] for n in names if n}
    for i, row in enumerate(rows):
        if i >= sample_rows:
            break
        for name, cell in zip(names, row, strict=False):
            if not name or cell is None:
                continue
            if hasattr(cell, "isoformat"):
                text = cell.isoformat() if hasattr(cell, "hour") else str(cell)
            else:
                text = str(cell)
            text = text.strip()
            if text:
                buckets[name].append(text)
    return [SourceColumn(name=name, samples=tuple(buckets.get(name, []))) for name in names if name]


def source_profile_from_scenario(scenario_dir: Path, meta: dict[str, Any]) -> SourceProfile:
    data_file = meta["data_file"]
    data_path = scenario_dir / data_file
    sample_rows = int(meta.get("sample_rows", 50))
    fmt = meta.get("data_format", data_path.suffix.lstrip(".").lower())
    if fmt == "csv":
        columns = _read_csv_columns(data_path, sample_rows)
    elif fmt in ("xlsx", "xlsm"):
        columns = _read_xlsx_columns(data_path, sample_rows)
    else:
        raise ValueError(f"unsupported data_format: {fmt}")
    return SourceProfile(columns=columns)


def target_schema_from_scenario(meta: dict[str, Any]) -> TargetSchema:
    fields = [
        TargetField(
            id=f["id"],
            labels=tuple(f.get("labels", [])),
            aliases=tuple(f.get("aliases", [])),
            expected_semantic_type=f.get("expected_semantic_type"),
            description=f.get("description"),
            exemplar_samples=tuple(f.get("exemplar_samples", [])),
        )
        for f in meta["target_fields"]
    ]
    return TargetSchema(fields=fields)
